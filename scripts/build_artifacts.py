from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, Frame, PageTemplate, PageBreak, Paragraph,
    Preformatted, Spacer, Table, TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
PPTX_OUT = ROOT / "apresentacao" / "apresentacao.pptx"
PDF_OUT = ROOT / "relatorio" / "relatorio.pdf"
REPORT_MD = ROOT / "relatorio" / "relatorio.md"

NAVY = RGBColor(18, 31, 53)
TEAL = RGBColor(31, 153, 148)
ORANGE = RGBColor(231, 137, 73)
WHITE = RGBColor(250, 252, 255)
LIGHT = RGBColor(230, 237, 243)
GRAY = RGBColor(87, 101, 116)


def set_bg(slide, color=NAVY):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_footer(slide, number):
    box = slide.shapes.add_textbox(Inches(.55), Inches(7.08), Inches(12.1), Inches(.2))
    p = box.text_frame.paragraphs[0]
    p.text = f"Pipeline ML + Otimização · Equipe Predictfy                                      {number}"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(159, 177, 194)


def title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(.65), Inches(.42), Inches(12), Inches(.75))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = WHITE
    if subtitle:
        sb = slide.shapes.add_textbox(Inches(.68), Inches(1.12), Inches(11.7), Inches(.45))
        sp = sb.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.name = "Arial"
        sp.font.size = Pt(13)
        sp.font.color.rgb = RGBColor(166, 220, 216)


def bullets(slide, items, x=.8, y=1.65, w=11.7, h=4.9, font=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            head, body = item
            p.text = f"{head}: {body}"
        else:
            p.text = item
        p.font.name = "Arial"
        p.font.size = Pt(font)
        p.font.color.rgb = WHITE
        p.space_after = Pt(14)
        p.level = 0
        p.text = "• " + p.text


def add_table(slide, data, widths=None, x=.65, y=1.65, w=12.0, h=4.8, font=13):
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    if widths:
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TEAL if r == 0 else (RGBColor(35, 51, 73) if r % 2 else RGBColor(43, 61, 84))
            cell.margin_left = Inches(.08)
            cell.margin_right = Inches(.08)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Arial"
                p.font.size = Pt(font)
                p.font.bold = r == 0
                p.font.color.rgb = WHITE
                p.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
    return table


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank); set_bg(slide)
    box = slide.shapes.add_textbox(Inches(.8), Inches(1.25), Inches(11.7), Inches(2.3))
    p = box.text_frame.paragraphs[0]
    p.text = "PIPELINE DE DECISÃO OPERACIONAL"
    p.font.name = "Arial"; p.font.bold = True; p.font.size = Pt(34); p.font.color.rgb = WHITE
    p2 = box.text_frame.add_paragraph(); p2.text = "Machine Learning, Otimização e Manutenção"; p2.font.name = "Arial"; p2.font.size = Pt(27); p2.font.color.rgb = RGBColor(166, 220, 216); p2.space_before = Pt(12)
    p3 = box.text_frame.add_paragraph(); p3.text = "Dados → ML → Cenários → Otimização → Decisão"; p3.font.name = "Arial"; p3.font.size = Pt(20); p3.font.color.rgb = ORANGE; p3.space_before = Pt(22)
    by = slide.shapes.add_textbox(Inches(.82), Inches(5.7), Inches(8), Inches(.6)); bp = by.text_frame.paragraphs[0]; bp.text = "Equipe Predictfy · FIAP"; bp.font.name = "Arial"; bp.font.size = Pt(17); bp.font.color.rgb = LIGHT

    slides = []
    slides.append(("Perguntas e resposta executiva", [
        ("Operação", "vazão ≈ 694,11 m³/h, temperatura ≈ 795,14 °C, pressão ≈ 33,32 bar e válvula ≈ 79,53%."),
        ("Manutenção", "inspeção imediata; executar se a degradação for confirmada."),
        ("Automação", "setpoints supervisionados; parada sempre aprovada por humano."),
        ("Limite", "o ganho econômico é um cenário didático, não um business case calibrado."),
    ]))
    slides.append(("Entendimento do negócio", [
        ("Controláveis", "Feedstock Flow, Reactor Temperature, Reactor Pressure e Valve Opening."),
        ("Estado", "Sensor Health, vibração e idade/tipo do catalisador."),
        ("Externas", "temperatura ambiente, hora, mês e unidade."),
        ("Restrições", "produção mínima, limites p1–p99 por unidade e suporte conjunto histórico."),
    ]))
    slides.append(("Dados e exploração", [
        "10.000 leituras a cada 4 horas, de 2020 a 2024, em três unidades.",
        "Sem nulos, duplicatas completas ou timestamps repetidos.",
        "Features temporais codificadas de forma cíclica.",
        "Vazão e saúde concentram a maior associação preditiva; importância não implica causalidade.",
    ]))

    for idx, (heading, items) in enumerate(slides, start=2):
        s = prs.slides.add_slide(blank); set_bg(s); title(s, heading); bullets(s, items); add_footer(s, idx)

    s = prs.slides.add_slide(blank); set_bg(s); title(s, "Target leakage: identidade exata", "Auditoria algébrica do target")
    bullets(s, [
        "Energy_Intensity = (3,6 × Electricity + 0,035 × Natural Gas) / Product Yield",
        "Erro máximo de reconstrução: 1,78 × 10⁻¹⁵.",
        "Electricity, Natural Gas e Yield não podem ser features pré-operação.",
        "Steam não entra na fórmula, mas também é medição pós-operação e foi excluído.",
    ], font=19); add_footer(s, 5)

    s = prs.slides.add_slide(blank); set_bg(s); title(s, "Modelos e validação temporal", "Walk-forward · calibração separada · teste final intocado")
    add_table(s, [
        ["Target", "Modelo", "R²", "MAE", "RMSE"],
        ["Energy", "Gradient Boosting", "0,6888", "0,2766", "0,3424"],
        ["Energy", "Híbrido físico ✓", "0,6930", "0,2761", "0,3401"],
        ["Yield", "Random Forest", "0,9999", "0,0126", "0,1250"],
        ["Yield", "Híbrido físico ✓", "1,0000", "0,0000", "0,0000"],
    ], widths=[2.2, 3.3, 1.5, 1.5, 1.5], h=3.3, font=14)
    bullets(s, ["RMSE 0,340; margem conformal 0,552; cobertura final 89,4%."], y=5.25, h=.8, font=17); add_footer(s, 6)

    s = prs.slides.add_slide(blank); set_bg(s); title(s, "Formulação da otimização")
    bullets(s, [
        "Minimizar EI prevista + desempate por menor movimento dos setpoints.",
        "Manter Yield previsto ≥ 64,37 ton/4h.",
        "Respeitar p1–p99 da unidade e proximidade a combinações observadas.",
        "Robusta: minimizar EI superior e exigir Yield inferior acima da meta.",
        "Validações: 10 mil candidatos aleatórios e linprog linear.",
    ], font=18); add_footer(s, 7)

    s = prs.slides.add_slide(blank); set_bg(s); title(s, "Caso de estudo sem seleção manual")
    add_table(s, [
        ["Estado", "Saúde", "Vibração", "Idade catalisador", "Origem"],
        ["Degradado", "0,578", "8,126 mm/s", "348 dias", "Maior score no teste"],
        ["Referência saudável", "0,970", "1,749 mm/s", "58 dias", "Linha real do treino"],
    ], widths=[2.2, 1.5, 2.0, 2.0, 3.2], h=2.5, font=14)
    bullets(s, ["Normalidade no treino: degradado 0,44º percentil; referência 20,11º percentil."], y=4.55, h=1.2, font=18); add_footer(s, 8)

    s = prs.slides.add_slide(blank); set_bg(s); title(s, "Solução operacional")
    add_table(s, [
        ["Variável", "Sem manutenção", "Manutenção imediata"],
        ["Feedstock Flow (m³/h)", "694,11", "694,11"],
        ["Reactor Temperature (°C)", "789,03", "795,14"],
        ["Reactor Pressure (bar)", "33,36", "33,32"],
        ["Valve Opening (%)", "79,91", "79,53"],
        ["Expected Yield (ton/4h)", "72,20", "121,25"],
        ["Energy Intensity", "3,106", "1,849"],
    ], widths=[4.3, 3.0, 3.5], h=4.7, font=14); add_footer(s, 9)

    s = prs.slides.add_slide(blank); set_bg(s); title(s, "Três cenários em 30 dias", "Custos parciais sob premissas didáticas")
    add_table(s, [
        ["Cenário", "Produção", "Energia", "Total parcial", "R$/ton"],
        ["Sem manutenção", "12.997 t", "40.350 un.", "R$ 5,436 mi", "418,27"],
        ["Postergada", "17.168 t", "39.902 un.", "R$ 5,322 mi", "309,97"],
        ["Imediata", "21.582 t", "39.902 un.", "R$ 5,261 mi", "243,76"],
    ], widths=[2.7, 2.1, 2.1, 2.7, 1.5], h=3.4, font=14)
    bullets(s, ["Economia parcial em 30 dias: R$ 175 mil; volumes ainda são diferentes."], y=5.3, h=.8, font=18); add_footer(s, 10)

    s = prs.slides.add_slide(blank); set_bg(s); title(s, "O que o resultado permite afirmar")
    bullets(s, [
        "Em 10 mil toneladas iguais: R$ 4,184 mi sem manter vs. R$ 2,462 mi com recuperação integral.",
        "Com recuperação nula, a manutenção perde R$ 45 mil; o efeito causal não foi observado.",
        "O ponto de equilíbrio da recuperação é aproximadamente 1,8%.",
        "Alternativas de produção entre 140% e 165% evitam operar sempre no extremo.",
        "Faltam receita, margem, custo completo da parada, estoque e demanda.",
        "EI nominal 1,849; limite superior conformal de 90% igual a 2,401.",
    ], font=19); add_footer(s, 11)

    s = prs.slides.add_slide(blank); set_bg(s); title(s, "Automação recomendada")
    add_table(s, [
        ["Decisão", "Nível", "Controle"],
        ["Qualidade, score e alerta", "Automático", "Monitoramento e drift"],
        ["Sugestão de setpoints", "Supervisionado", "Limites oficiais + rollback"],
        ["Parada para manutenção", "Aprovação humana", "Engenharia, segurança e janela"],
    ], widths=[3.8, 2.5, 5.0], h=3.6, font=14); add_footer(s, 12)

    s = prs.slides.add_slide(blank); set_bg(s); title(s, "Decisão final")
    bullets(s, [
        ("Configuração", "Flow 694,11 · Temp 795,14 · Pressure 33,32 · Valve 79,53."),
        ("Produção/energia", "121,25 ton/4h e EI 1,849 (RMSE 0,340)."),
        ("Manutenção", "sim, condicionada à confirmação por inspeção humana."),
        ("Total parcial", "R$ 5,261 milhões em 30 dias, sob premissas."),
    ], font=19); add_footer(s, 13)

    s = prs.slides.add_slide(blank); set_bg(s); title(s, "Conclusão e próximos dados")
    bullets(s, [
        "Recomendação operacional dentro do suporte histórico, sujeita aos limites oficiais.",
        "Automatizar monitoramento e sugestão; manter a decisão de parada com humano.",
        "Para produção: falhas, ordens de manutenção, antes/depois, limites oficiais, preços e margens.",
        "Também são necessários demanda, estoque, duração real de parada e penalidades.",
    ], font=19); add_footer(s, 14)

    prs.save(PPTX_OUT)


def inline_markup(text):
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"`(.+?)`", r'<font name="ArialUnicode">\1</font>', text)
    return text


def build_pdf():
    pdfmetrics.registerFont(TTFont("ArialUnicode", "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"))
    pdfmetrics.registerFont(TTFont("Arial", "/System/Library/Fonts/Supplemental/Arial.ttf"))
    pdfmetrics.registerFont(TTFont("ArialBold", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"))

    styles = getSampleStyleSheet()
    normal = ParagraphStyle("Body", parent=styles["BodyText"], fontName="Arial", fontSize=9.2, leading=12.5, spaceAfter=6, textColor=colors.HexColor("#1f2933"))
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontName="ArialBold", fontSize=18, leading=22, spaceAfter=12, textColor=colors.HexColor("#122035"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontName="ArialBold", fontSize=13, leading=16, spaceBefore=8, spaceAfter=7, textColor=colors.HexColor("#1f9994"))
    h3 = ParagraphStyle("H3", parent=styles["Heading3"], fontName="ArialBold", fontSize=10.5, leading=13, spaceBefore=6, spaceAfter=5, textColor=colors.HexColor("#455a6d"))
    bullet = ParagraphStyle("Bullet", parent=normal, leftIndent=12, firstLineIndent=-7, bulletIndent=3)
    code_style = ParagraphStyle("Code", parent=normal, fontName="ArialUnicode", fontSize=8, leading=10, backColor=colors.HexColor("#eef2f5"), leftIndent=8, rightIndent=8, borderPadding=6)

    def header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Arial", 7.5)
        canvas.setFillColor(colors.HexColor("#657789"))
        canvas.drawString(1.6 * cm, A4[1] - 1.0 * cm, "Pipeline de ML, Otimização e Decisão de Manutenção")
        canvas.drawRightString(A4[0] - 1.6 * cm, .85 * cm, f"Página {doc.page}")
        canvas.restoreState()

    doc = BaseDocTemplate(str(PDF_OUT), pagesize=A4, leftMargin=1.6*cm, rightMargin=1.6*cm, topMargin=1.6*cm, bottomMargin=1.4*cm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates(PageTemplate(id="main", frames=frame, onPage=header_footer))

    lines = REPORT_MD.read_text(encoding="utf-8").splitlines()
    story, i = [], 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            story.append(Spacer(1, 3)); i += 1; continue
        if line.startswith("```"):
            block = []; i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i]); i += 1
            story.append(Preformatted("\n".join(block), code_style)); i += 1; continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i]); i += 1
            rows = [[inline_markup(c.strip()) for c in row.strip("|").split("|")] for row in table_lines]
            rows = [rows[0]] + [r for r in rows[2:]]
            pdata = [[Paragraph(c, normal) for c in row] for row in rows]
            col_width = doc.width / len(pdata[0])
            table = Table(pdata, colWidths=[col_width] * len(pdata[0]), repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1f9994")),
                ("TEXTCOLOR", (0,0), (-1,0), colors.white),
                ("FONTNAME", (0,0), (-1,0), "ArialBold"),
                ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#aebbc6")),
                ("VALIGN", (0,0), (-1,-1), "TOP"),
                ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f2f6f8")]),
                ("LEFTPADDING", (0,0), (-1,-1), 4), ("RIGHTPADDING", (0,0), (-1,-1), 4),
                ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
            ]))
            story.extend([table, Spacer(1, 7)]); continue
        if line.startswith("### "):
            story.append(Paragraph(inline_markup(line[4:]), h3)); i += 1; continue
        if line.startswith("## "):
            story.append(Paragraph(inline_markup(line[3:]), h2)); i += 1; continue
        if line.startswith("# "):
            story.append(Paragraph(inline_markup(line[2:]), h1)); i += 1; continue
        if line.startswith("- "):
            story.append(Paragraph("• " + inline_markup(line[2:]), bullet)); i += 1; continue
        paragraph = [line]; i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "- ", "```")):
            paragraph.append(lines[i].strip()); i += 1
        story.append(Paragraph(inline_markup(" ".join(paragraph)), normal))

    doc.build(story)


if __name__ == "__main__":
    build_pdf()
    print(f"PDF: {PDF_OUT}")
