"""Gera um PPTX visualmente idêntico ao PDF revisado da apresentação."""

from pathlib import Path
import shutil
import subprocess
import tempfile

from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / "apresentacao" / "apresentacao.pdf"
OUT = ROOT / "apresentacao" / "apresentacao.pptx"


def build() -> None:
    if not PDF.exists():
        raise FileNotFoundError(f"Renderize primeiro o PDF: {PDF}")
    renderer = shutil.which("pdftoppm")
    if renderer is None:
        raise RuntimeError("pdftoppm não foi localizado no PATH")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory(prefix="petro-slides-") as tmp:
        prefix = Path(tmp) / "slide"
        subprocess.run(
            [renderer, "-png", "-r", "144", str(PDF), str(prefix)],
            check=True,
            stdout=subprocess.DEVNULL,
        )
        pages = sorted(Path(tmp).glob("slide-*.png"))
        if len(pages) != 16:
            raise RuntimeError(f"Esperadas 16 páginas; encontradas {len(pages)}")
        for page in pages:
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(page), 0, 0, width=prs.slide_width, height=prs.slide_height
            )

    prs.core_properties.title = "Pipeline Petroquímico — ML, Otimização e Decisão"
    prs.core_properties.subject = "Apresentação final — FIAP"
    prs.core_properties.author = "Equipe Predictfy"
    prs.save(OUT)
    print(f"PPTX criado com {len(prs.slides)} slides: {OUT}")


if __name__ == "__main__":
    build()
