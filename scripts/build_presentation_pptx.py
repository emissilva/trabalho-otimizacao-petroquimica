from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import os
import shutil
import tempfile
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "apresentacao" / "apresentacao.pptx"
FONT_DIR = ROOT / "fontes"

PAPER = RGBColor(21, 23, 27)
INK = RGBColor(244, 245, 247)
MUTED = RGBColor(155, 163, 175)
CORAL = RGBColor(255, 91, 107)
TEAL = RGBColor(23, 195, 210)
SUN = RGBColor(246, 196, 83)
LILAC = RGBColor(40, 38, 59)
MINT = RGBColor(23, 56, 59)
WHITE = RGBColor(34, 38, 44)
LINE = RGBColor(52, 57, 65)
FONT = "Inter"


def rgb(hex_value):
    hex_value = hex_value.lstrip("#")
    return RGBColor.from_string(hex_value.upper())


def shape(slide, kind, x, y, w, h, fill, line=None, radius=False):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line or fill
    return shp


def textbox(slide, text, x, y, w, h, size=18, color=INK, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    return box


def richbox(slide, runs, x, y, w, h, size=18, align=PP_ALIGN.LEFT, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.05); tf.margin_right = Inches(.05)
    p = tf.paragraphs[0]; p.alignment = align
    for text, run_color, bold in runs:
        r = p.add_run(); r.text = text; r.font.name = FONT; r.font.size = Pt(size)
        r.font.bold = bold; r.font.color.rgb = run_color or color
    return box


def base(prs, number, kicker, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = PAPER
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, .025, CORAL, CORAL)
    textbox(slide, kicker.upper(), .75, .45, 5.8, .24, 9, CORAL, True)
    textbox(slide, title, .75, .76, 11.3, .9, 28, INK, True)
    if subtitle:
        textbox(slide, subtitle, .78, 1.58, 9.5, .42, 12, MUTED)
    shape(slide, MSO_SHAPE.RECTANGLE, .75, 6.96, 4.0, .02, CORAL, CORAL)
    textbox(slide, "PIPELINE PETROQUÍMICO", .75, 7.18, 4.0, .15, 7.5, MUTED, True)
    textbox(slide, f"{number:02d}", 12.1, 7.16, .45, .18, 8, INK, True, PP_ALIGN.RIGHT)
    return slide


def pill(slide, text, x, y, w, color=INK, fill=WHITE):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, .34, fill, fill)
    textbox(slide, text, x + .08, y + .05, w - .16, .22, 9, color, True, PP_ALIGN.CENTER)


def metric(slide, x, y, w, label, value, desc, accent):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, 1.3, WHITE, LINE)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, .07, 1.3, accent, accent)
    textbox(slide, label.upper(), x + .22, y + .18, w - .4, .18, 8, MUTED, True)
    textbox(slide, value, x + .22, y + .42, w - .4, .43, 24, INK, True)
    textbox(slide, desc, x + .22, y + .92, w - .4, .25, 9, MUTED)


def footer_note(slide, text, x=.75, y=6.75, w=11.0):
    textbox(slide, text, x, y, w, .2, 9, MUTED)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 — cover
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = PAPER
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, .025, CORAL, CORAL)
    textbox(s, "FIAP · MACHINE LEARNING + OTIMIZAÇÃO", .75, 1.08, 6, .25, 10, CORAL, True)
    textbox(s, "Dados viram decisão.\nCom limites claros.", .75, 1.48, 7.7, 1.45, 38, INK, True)
    shape(s, MSO_SHAPE.RECTANGLE, .75, 3.25, .82, .04, CORAL, CORAL)
    textbox(s, "Otimização petroquímica para equilibrar produção, energia e manutenção — sem esconder incerteza.", .75, 3.55, 6.8, .75, 16, MUTED)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 9.25, 1.82, 3.0, 2.05, WHITE, LINE)
    shape(s, MSO_SHAPE.RECTANGLE, 9.25, 1.82, .04, 2.05, CORAL, CORAL)
    textbox(s, "14", 9.65, 2.15, 1.7, .55, 30, INK, True)
    textbox(s, "slides · 11–12 min\n\nKelvin Douglas Rabelo", 9.65, 2.78, 2.15, .72, 10, MUTED, True)
    shape(s, MSO_SHAPE.RECTANGLE, .75, 6.96, 4.0, .02, CORAL, CORAL)
    textbox(s, "COMPLEXO PETROQUÍMICO · ENTREGA FINAL · 2026", .75, 7.16, 5, .18, 8, MUTED, True)
    textbox(s, "01", 12.1, 7.16, .45, .18, 8, INK, True, PP_ALIGN.RIGHT)

    # 2 — executive answer
    s = base(prs, 2, "01 · resposta executiva", "A recomendação cabe em três movimentos.")
    metric(s, .75, 2.18, 3.72, "Configuração", "694,11", "m³/h · 795,14 °C · 33,32 bar", TEAL)
    metric(s, 4.8, 2.18, 3.72, "Capacidade", "121,25", "toneladas por intervalo de 4 horas", CORAL)
    metric(s, 8.85, 2.18, 3.72, "Decisão", "1,8%", "recuperação mínima para empatar", SUN)
    shape(s, MSO_SHAPE.RECTANGLE, .75, 4.65, 11.82, .03, INK, INK)
    richbox(s, [("A máquina pode ", INK, True), ("sugerir", CORAL, True), (" setpoints. A planta ainda precisa de uma pessoa para autorizar a parada.", INK, True)], .75, 5.0, 11.5, .7, 22)

    # 3 — business flow
    s = base(prs, 3, "02 · entendimento do negócio", "O problema não começa no solver.", "Separamos controle, estado, decisão e resultado antes de otimizar.")
    labels = [("Dados", "10.000 leituras", WHITE), ("Estado", "saúde · vibração", LILAC), ("Decisão", "flow · temp · pressão", WHITE), ("Resultado", "yield · energia", LILAC)]
    x = .85
    for i, (head, sub, fill) in enumerate(labels):
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.0, 2.3, 1.0, fill, LINE)
        textbox(s, head, x + .1, 3.25, 2.1, .25, 18, INK, True, PP_ALIGN.CENTER)
        textbox(s, sub, x + .1, 3.62, 2.1, .18, 9, MUTED, False, PP_ALIGN.CENTER)
        if i < 3: textbox(s, "→", x + 2.43, 3.27, .55, .3, 25, CORAL, True, PP_ALIGN.CENTER)
        x += 3.05
    footer_note(s, "histórico observado                    condição do ativo                    variáveis ajustáveis                    targets pós-operação")

    # 4 — EDA
    s = base(prs, 4, "03 · exploração e feature engineering", "A base é sintética — e isso também é um achado.")
    metric(s, .75, 2.1, 2.75, "Observações", "10k", "leituras a cada 4 horas", LILAC)
    metric(s, 3.72, 2.1, 2.75, "Período", "2020–24", "janela temporal", TEAL)
    metric(s, 6.69, 2.1, 2.75, "Unidades", "3", "plantas produtivas", SUN)
    metric(s, 9.66, 2.1, 2.75, "Qualidade", "0", "nulos e duplicatas", CORAL)
    shape(s, MSO_SHAPE.RECTANGLE, .75, 4.35, 5.8, .03, INK, INK)
    textbox(s, "Yield = 0,18 × Flow × Health", .75, 4.6, 5.8, .5, 22, INK, True)
    textbox(s, "Identidade exata com informação pré-operação. Não é leakage; é a estrutura geradora da base.", .75, 5.23, 5.6, .48, 12, MUTED)
    # micro bars
    textbox(s, "R² no teste final", 7.1, 4.35, 3, .2, 10, MUTED, True)
    for i, (label, val, col) in enumerate([("yield híbrido", 1.0, TEAL), ("energia híbrida", .693, CORAL)]):
        y = 4.82 + i * .62
        textbox(s, label, 7.1, y, 1.55, .2, 10, INK, True)
        shape(s, MSO_SHAPE.RECTANGLE, 8.75, y + .04, 3.1, .18, LINE, LINE)
        shape(s, MSO_SHAPE.RECTANGLE, 8.75, y + .04, 3.1 * val, .18, col, col)
        textbox(s, f"{val:.3f}", 12.0, y - .01, .45, .2, 10, INK, True, PP_ALIGN.RIGHT)

    # 5 — leakage
    s = base(prs, 5, "04 · auditoria de leakage", "A resposta não pode estar escondida na pergunta.")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, .75, 2.25, 11.8, 1.85, INK, INK)
    richbox(s, [("Energy Intensity", SUN, True), (" = (3,6 × ", WHITE, True), ("Electricity", CORAL, True), (" + 0,035 × ", WHITE, True), ("Natural Gas", CORAL, True), (") / ", WHITE, True), ("Yield", SUN, True)], 1.1, 2.85, 11.0, .48, 24)
    textbox(s, "Eletricidade, gás e produção só aparecem depois da operação. Entram na auditoria e como targets auxiliares — nunca como features pré-operacionais.", 1.1, 3.55, 10.6, .3, 11, RGBColor(189,197,213))
    shape(s, MSO_SHAPE.RECTANGLE, .75, 4.62, 5.68, .78, MINT, TEAL)
    textbox(s, "SEGURO", .98, 4.84, 1.0, .16, 9, TEAL, True)
    textbox(s, "flow · health · vibração · catalisador · ambiente · tempo", 2.0, 4.79, 4.15, .25, 11, INK, True)
    shape(s, MSO_SHAPE.RECTANGLE, 6.87, 4.62, 5.68, .78, WHITE, CORAL)
    textbox(s, "EXCLUÍDO", 7.1, 4.84, 1.12, .16, 9, CORAL, True)
    textbox(s, "electricity · gás · vapor · medições pós-operação", 8.28, 4.79, 3.95, .25, 11, INK, True)

    # 6 — models
    s = base(prs, 6, "05 · modelos e validação", "A métrica melhorou; a confiança ficou explícita.")
    textbox(s, "RMSE no teste final", .85, 2.15, 4, .25, 11, MUTED, True)
    bars = [("EI · híbrido", .340, CORAL), ("EI · boosting", .342, SUN), ("Yield · híbrido", 0.0, TEAL)]
    for i, (lab, val, col) in enumerate(bars):
        x = 1.0 + i * 3.75
        h = 2.15 if val == 0 else 2.15 * (1 - val / .4)
        h = max(.18, h)
        shape(s, MSO_SHAPE.RECTANGLE, x, 5.2 - h, 1.05, h, col, col)
        textbox(s, f"{val:.3f}", x - .1, 5.2 - h - .32, 1.25, .22, 13, INK, True, PP_ALIGN.CENTER)
        textbox(s, lab, x - .45, 5.36, 1.95, .35, 10, MUTED, True, PP_ALIGN.CENTER)
    shape(s, MSO_SHAPE.RECTANGLE, .85, 5.2, 10.9, .025, INK, INK)
    pill(s, "walk-forward estável", .85, 6.0, 1.75, TEAL, MINT)
    pill(s, "calibração conformal", 2.82, 6.0, 1.85, CORAL, WHITE)
    textbox(s, "89,4% de cobertura no teste final", 4.93, 6.05, 3.5, .2, 11, MUTED, True)

    # 7 — optimization
    s = base(prs, 7, "06 · otimização", "O solver procura o melhor ponto dentro do mundo observado.")
    textbox(s, "min EI(x) + movimento", .85, 2.22, 4.9, .45, 24, INK, True)
    for i, (lab, value, col, right) in enumerate([("Yield", .72, TEAL, "≥ meta"), ("Suporte", .51, CORAL, "≤ p99"), ("Robustez", .90, SUN, "EI +0,552")]):
        y = 3.0 + i * .62
        textbox(s, lab, .9, y, 1.0, .2, 11, INK, True)
        shape(s, MSO_SHAPE.RECTANGLE, 2.0, y + .04, 2.8, .17, LINE, LINE)
        shape(s, MSO_SHAPE.RECTANGLE, 2.0, y + .04, 2.8 * value, .17, col, col)
        textbox(s, right, 4.98, y, 1.1, .2, 10, MUTED, True)
    textbox(s, "A distância condicionada combina setpoints + estado do ativo.", .85, 5.2, 5.2, .28, 11, MUTED)
    shape(s, MSO_SHAPE.RECTANGLE, 7.0, 2.35, 4.5, 2.7, WHITE, LINE)
    # simple support curve
    textbox(s, "suporte histórico", 7.35, 2.65, 2.1, .2, 10, MUTED, True)
    shape(s, MSO_SHAPE.RECTANGLE, 7.55, 4.45, 3.5, .02, INK, INK)
    shape(s, MSO_SHAPE.RECTANGLE, 7.55, 3.0, .02, 1.45, INK, INK)
    for cx, cy, col in [(8.15,4.1,CORAL),(8.55,3.75,CORAL),(9.2,3.4,TEAL),(9.95,3.0,SUN),(10.55,2.7,SUN)]:
        shape(s, MSO_SHAPE.OVAL, cx, cy, .15, .15, col, col)
    textbox(s, "atual →", 7.55, 4.58, 1, .2, 9, MUTED)
    textbox(s, "ótimo →", 10.25, 2.43, 1, .2, 9, MUTED)

    # 8 — asset state
    s = base(prs, 8, "07 · estado do ativo", "A manutenção muda o estado — não inventa uma observação.")
    shape(s, MSO_SHAPE.RECTANGLE, 1.0, 3.25, 10.55, .05, INK, INK)
    for i, (x, head, body, col) in enumerate([(1.0,"Atual","saúde 0,578\nvibração 8,126\ncatalisador 348d",CORAL),(5.15,"Inspeção","confirmar sensores\nlimites de engenharia\njanela de parada",SUN),(9.3,"Referência","saúde 0,970\nvibração 1,749\ncatalisador 58d",TEAL)]):
        shape(s, MSO_SHAPE.OVAL, x, 3.0, .55, .55, col, PAPER)
        textbox(s, head, x, 3.82, 2.3, .28, 19, INK, True)
        textbox(s, body, x, 4.25, 2.4, .75, 12, MUTED)
    shape(s, MSO_SHAPE.OVAL, 10.45, 5.25, 1.35, 1.35, LILAC, LILAC)
    textbox(s, "1,8%", 10.55, 5.55, 1.15, .25, 20, INK, True, PP_ALIGN.CENTER)
    textbox(s, "recuperação mínima", 10.35, 5.88, 1.55, .3, 8, INK, True, PP_ALIGN.CENTER)

    # 9 — setpoints
    s = base(prs, 9, "08 · configuração recomendada", "Uma configuração com margem visível.")
    values = [("Flow", "694,11", .88, TEAL), ("Temperature", "795,14 °C", .82, CORAL), ("Pressure", "33,32 bar", .62, SUN), ("Valve", "79,53%", .84, rgb("988CFF"))]
    for i, (lab, val, frac, col) in enumerate(values):
        x = .85 if i < 2 else 6.9; y = 2.35 + (i % 2) * .95
        textbox(s, lab, x, y, 1.2, .2, 12, INK, True)
        shape(s, MSO_SHAPE.RECTANGLE, x + 1.35, y + .03, 2.65, .16, LINE, LINE)
        shape(s, MSO_SHAPE.RECTANGLE, x + 1.35, y + .03, 2.65 * frac, .16, col, col)
        textbox(s, val, x + 4.2, y - .03, 1.2, .22, 12, INK, True)
    shape(s, MSO_SHAPE.RECTANGLE, .85, 5.0, 11.6, .85, INK, INK)
    textbox(s, "121,25 ton/4h", 1.15, 5.19, 3.0, .3, 22, WHITE, True)
    textbox(s, "EI 1,849  ·  limite superior robusto 2,401", 8.15, 5.2, 3.8, .25, 13, SUN, True, PP_ALIGN.RIGHT)
    textbox(s, "produção prevista · suporte condicionado 1,035 / 2,032", 1.15, 5.57, 5.0, .16, 8, RGBColor(200,210,222))

    # 10 — scenarios
    s = base(prs, 10, "09 · três cenários", "A economia depende da base de comparação.")
    headers = ["Cenário", "Produção", "Energia", "Total parcial", "R$/ton"]
    rows = [["Sem manutenção","12.997 t","40.350","R$ 5,436 mi","418,27"],["Postergada","17.168 t","39.902","R$ 5,322 mi","309,97"],["Imediata","21.582 t","39.902","R$ 5,261 mi","243,76"]]
    x0,y0 = .85,2.3; widths=[2.2,1.65,1.65,2.25,1.4]
    x=x0
    for h,w in zip(headers,widths): shape(s,MSO_SHAPE.RECTANGLE,x,y0,w,.48,INK,INK); textbox(s,h,x+.08,y0+.14,w-.16,.16,9,WHITE,True); x+=w
    for ri,row in enumerate(rows):
        x=x0; fill=MINT if ri==2 else WHITE
        for val,w in zip(row,widths): shape(s,MSO_SHAPE.RECTANGLE,x,y0+.48+ri*.54,w,.54,fill,LINE); textbox(s,val,x+.08,y0+.67+ri*.54,w-.16,.16,10,TEAL if ri==2 else INK,ri==2); x+=w
    textbox(s, "Em 30 dias: R$ 175 mil de economia parcial", .85, 5.2, 5.8, .35, 20, CORAL, True)
    textbox(s, "Volumes são diferentes; por isso também comparamos uma meta comum de 10 mil toneladas.", .85, 5.72, 7.6, .3, 11, MUTED)

    # 11 — robust
    s = base(prs, 11, "10 · decisão robusta", "A margem não muda o ponto — muda o que podemos prometer.")
    textbox(s, "EI nominal", .9, 2.28, 2, .22, 12, MUTED, True)
    textbox(s, "1,849", .9, 2.62, 2.2, .5, 32, CORAL, True)
    shape(s, MSO_SHAPE.RECTANGLE, 1.0, 4.1, 4.5, .1, LILAC, LILAC); shape(s,MSO_SHAPE.RECTANGLE,2.1,4.1,2.55,.1,TEAL,TEAL)
    for x,col in [(1.0,SUN),(3.55,CORAL),(4.65,TEAL)]: shape(s,MSO_SHAPE.OVAL,x,3.9,.5,.5,col,PAPER)
    textbox(s,"limite inferior",.85,4.55,1.25,.18,8,MUTED); textbox(s,"nominal",3.35,4.55,1.0,.18,8,MUTED); textbox(s,"limite 2,401",4.45,4.55,1.1,.18,8,MUTED)
    textbox(s, "Yield nominal", 7.0, 2.28, 2.2, .22, 12, MUTED, True)
    textbox(s, "121,25", 7.0, 2.62, 3.0, .5, 32, TEAL, True)
    shape(s, MSO_SHAPE.RECTANGLE, 7.0, 3.9, 4.9, 1.1, MINT, MINT)
    textbox(s, "limite inferior = 121,25", 7.3, 4.18, 4.2, .26, 18, INK, True)
    textbox(s, "erro da identidade: 1,42 × 10⁻¹⁴", 7.3, 4.58, 4.0, .18, 10, MUTED)
    textbox(s, "A produção tem garantia algébrica na base; a energia precisa ser comunicada como intervalo.", .9, 5.65, 10.7, .3, 14, INK, True)

    # 12 — alternatives
    s = base(prs, 12, "11 · alternativas operacionais", "Nem sempre o melhor ponto é o ponto certo.")
    opts=[("140%","91,92 t","2,439","526,21",TEAL),("150%","98,48 t","2,276","563,79",SUN),("160%","105,05 t","2,134","601,38",CORAL),("165%","108,33 t","2,069","620,17",rgb("988CFF"))]
    for i,(meta,yld,ei,flow,col) in enumerate(opts):
        x=.85+i*3.02; shape(s,MSO_SHAPE.RECTANGLE,x,2.45,2.65,2.05,WHITE,LINE); shape(s,MSO_SHAPE.RECTANGLE,x,2.45,2.65,.09,col,col)
        textbox(s,f"META {meta}",x+.2,2.78,2.2,.2,10,MUTED,True); textbox(s,yld,x+.2,3.18,2.2,.42,25,INK,True); textbox(s,f"EI {ei}",x+.2,3.85,1.2,.2,13,col,True); textbox(s,f"Flow {flow}",x+.2,4.15,1.8,.18,10,MUTED)
    textbox(s,"Escolher uma meta de capacidade também é uma decisão de negócio.", .85, 5.55, 9.5, .35, 21, INK, True)

    # 13 — automation
    s = base(prs, 13, "12 · automação", "Automatizar o fluxo, não terceirizar o julgamento.")
    steps=[("NÍVEL 1","Automático","dados · scores · drift",WHITE,140),("NÍVEL 2","Supervisionado","setpoints · limites · rollback",MINT,190),("NÍVEL 3","Aprovação humana","parada · segurança · janela",RGBColor(42,32,36),245)]
    for i,(level,head,desc,fill,height) in enumerate(steps):
        x=.95+i*4.1; y=5.45-height/100
        shape(s,MSO_SHAPE.RECTANGLE,x,y,3.45,height/100,fill,TEAL if i==1 else fill)
        textbox(s,level,x+.22,y+.22,2.8,.18,9,CORAL if i<2 else SUN,True); textbox(s,head,x+.22,y+.56,2.8,.3,18,INK,True); textbox(s,desc,x+.22,y+1.0,2.75,.45,11,MUTED)

    # 14 — closing
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = PAPER
    shape(s,MSO_SHAPE.OVAL,10.35,-.65,4.2,4.2,LILAC,LILAC); shape(s,MSO_SHAPE.OVAL,10.6,6.25,1.2,1.2,SUN,SUN)
    textbox(s,"13 · FECHAMENTO",.75,1.0,4,.22,10,CORAL,True); textbox(s,"A recomendação ficou mais honesta —\ne mais útil.",.75,1.4,8.8,1.15,34,INK,True); shape(s,MSO_SHAPE.RECTANGLE,.75,2.95,1.7,.06,TEAL,TEAL)
    closing=["Produção e energia foram separadas.","A incerteza agora aparece na decisão.","A manutenção tem limiar, não promessa.","O humano continua no ponto irreversível."]
    for i,text in enumerate(closing):
        x=.85+(i%2)*5.3; y=3.55+(i//2)*.65
        textbox(s,"0"+str(i+1),x,y, .35,.2,10,TEAL,True); textbox(s,text,x+.48,y-.02,4.4,.26,14,INK,True)
    textbox(s,"Próximo passo para produção: falhas reais, ordens de manutenção, limites oficiais, preços e margem.",.75,5.45,8.0,.4,14,MUTED)
    textbox(s,"OBRIGADO",.75,7.16,2,.18,8,MUTED,True); textbox(s,"14",12.1,7.16,.45,.18,8,INK,True,PP_ALIGN.RIGHT)
    prs.save(OUT)
    embed_fonts(OUT)
    print(f"PPTX criado: {OUT}")


def embed_fonts(pptx_path):
    """Add Inter font parts and OOXML relationships to the generated package."""
    ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    rel_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
    ET.register_namespace("p", ns)
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        temp_path = Path(tmp.name)
    try:
        with ZipFile(pptx_path, "r") as zin, ZipFile(temp_path, "w", ZIP_DEFLATED) as zout:
            presentation = ET.fromstring(zin.read("ppt/presentation.xml"))
            relationships = ET.fromstring(zin.read("ppt/_rels/presentation.xml.rels"))
            content_types = ET.fromstring(zin.read("[Content_Types].xml"))
            ids = [int(el.attrib["Id"][3:]) for el in relationships if el.attrib.get("Id", "").startswith("rId")]
            next_id = max(ids, default=0) + 1
            embedded = presentation.find(f"{{{ns}}}embeddedFontLst")
            if embedded is None:
                embedded = ET.SubElement(presentation, f"{{{ns}}}embeddedFontLst")
            font = ET.SubElement(embedded, f"{{{ns}}}embeddedFont")
            ET.SubElement(font, f"{{{ns}}}font", {"typeface": FONT})
            for tag, filename in [("regular", "Inter-Regular.ttf"), ("bold", "Inter-Bold.ttf")]:
                rid = f"rId{next_id}"; next_id += 1
                ET.SubElement(font, f"{{{ns}}}{tag}", {"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": rid})
                ET.SubElement(relationships, f"{{{rel_ns}}}Relationship", {
                    "Id": rid, "Type": rel_type, "Target": f"fonts/{filename}"
                })
            ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
            if not any(el.attrib.get("Extension") == "ttf" for el in content_types):
                ET.SubElement(content_types, f"{{{ct_ns}}}Default", {
                    "Extension": "ttf", "ContentType": "application/x-font-ttf"
                })
            replacements = {
                "ppt/presentation.xml": ET.tostring(presentation, encoding="utf-8", xml_declaration=True),
                "ppt/_rels/presentation.xml.rels": ET.tostring(relationships, encoding="utf-8", xml_declaration=True),
                "[Content_Types].xml": ET.tostring(content_types, encoding="utf-8", xml_declaration=True),
            }
            for item in zin.infolist():
                if item.filename in replacements or item.filename.startswith("ppt/fonts/"):
                    continue
                zout.writestr(item, zin.read(item.filename))
            for name, data in replacements.items(): zout.writestr(name, data)
            zout.writestr("ppt/fonts/Inter-Regular.ttf", (FONT_DIR / "Inter-Regular.ttf").read_bytes())
            zout.writestr("ppt/fonts/Inter-Bold.ttf", (FONT_DIR / "Inter-Bold.ttf").read_bytes())
        os.replace(temp_path, pptx_path)
    finally:
        if temp_path.exists(): temp_path.unlink()


if __name__ == "__main__":
    build()
